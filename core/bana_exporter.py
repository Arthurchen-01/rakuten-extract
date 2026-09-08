# -*- coding: utf-8 -*-
"""
Bana 16-Column Official Benchmark Excel Exporter
法定 16 列 Bana 标准真本工作簿导出器
具备 4 大工作表视图守恒与自适应排版防挤压机制
"""
import os
import sys
import json
import hashlib
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

STANDARD_HEADERS_16 = [
    "账号", "密码", "姓", "名", "姓（假名）", "名（假名）",
    "其他信息", "生日", "性别", "地址",
    "店铺名称1", "订单1",
    "店铺名称2", "订单2",
    "店铺名称3", "订单3"
]

def sha256_file(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest().upper()

def export_bana_excel(data_source, output_excel_path, total_accounts=None):
    """
    将萃取数据导出为法定 16 列 Bana 标准工作簿。
    支持输入 list 字典、单个 json 文件路径或包含 result_*.json 的目录路径。
    """
    all_accounts = []
    if isinstance(data_source, list):
        all_accounts = data_source
    elif isinstance(data_source, (str, Path)):
        p = Path(data_source)
        if p.is_dir():
            for jf in sorted(list(p.glob("result_*.json"))):
                try:
                    with open(jf, 'r', encoding='utf-8') as f:
                        all_accounts.append(json.load(f))
                except Exception:
                    pass
        elif p.is_file():
            with open(p, 'r', encoding='utf-8') as f:
                content = json.load(f)
                all_accounts = content if isinstance(content, list) else [content]
                
    out_path = Path(output_excel_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    
    wb = openpyxl.Workbook()
    
    # Sheet 1: 乐天导出导入 (法定 16 列主表)
    ws1 = wb.active
    ws1.title = "乐天导出导入"
    
    # Sheet 2: 口径说明
    ws2 = wb.create_sheet(title="口径说明")
    
    # Sheet 3: 账号检测与现场快照总览
    ws3 = wb.create_sheet(title="账号检测与现场快照总览")
    
    # Sheet 4: 订单明细总览
    ws4 = wb.create_sheet(title="订单明细总览")
    
    header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="BF0000", end_color="BF0000", fill_type="solid") # 乐天法定红
    sub_header_fill = PatternFill(start_color="2A3F54", end_color="2A3F54", fill_type="solid") # 深海蓝
    
    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="top", wrap_text=True)
    
    # --- 编写 Sheet 1: 乐天导出导入 ---
    ws1.row_dimensions[1].height = 28
    for col_idx, h in enumerate(STANDARD_HEADERS_16, 1):
        cell = ws1.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border
        
    for row_counter, acc in enumerate(all_accounts, 2):
        prof = acc.get('profile', {})
        bana_orders = acc.get('bana_orders', [])
        raw_orders = acc.get('raw_orders', [])
        
        email = acc.get('email', '')
        pwd = acc.get('password', '')
        last_name = prof.get('last_name', '')
        first_name = prof.get('first_name', '')
        last_kana = prof.get('last_name_kana', '')
        first_kana = prof.get('first_name_kana', '')
        card_info = prof.get('card_info', '')
        bday = prof.get('birthday', '')
        gender = prof.get('gender', '')
        addr = prof.get('address', '')
        
        s1 = raw_orders[0].get('shop_name', '') if len(raw_orders) > 0 else ''
        o1 = bana_orders[0] if len(bana_orders) > 0 else ''
        s2 = raw_orders[1].get('shop_name', '') if len(raw_orders) > 1 else ''
        o2 = bana_orders[1] if len(bana_orders) > 1 else ''
        s3 = raw_orders[2].get('shop_name', '') if len(raw_orders) > 2 else ''
        o3 = bana_orders[2] if len(bana_orders) > 2 else ''
        
        row_vals = [
            email, pwd, last_name, first_name, last_kana, first_kana,
            card_info, bday, gender, addr,
            s1, o1, s2, o2, s3, o3
        ]
        
        ws1.row_dimensions[row_counter].height = 95 if (addr or o1) else 25
        for col_idx, val in enumerate(row_vals, 1):
            cell = ws1.cell(row=row_counter, column=col_idx, value=val)
            cell.font = Font(name="微软雅黑", size=10)
            cell.alignment = align_left
            cell.border = thin_border
            
    col_widths_16 = {
        1: 28, 2: 16, 3: 10, 4: 10, 5: 12, 6: 12,
        7: 20, 8: 12, 9: 8, 10: 38,
        11: 22, 12: 45, 13: 22, 14: 45, 15: 22, 16: 45
    }
    for col_idx, w in col_widths_16.items():
        ws1.column_dimensions[get_column_letter(col_idx)].width = w

    # --- 编写 Sheet 2: 口径说明 ---
    ws2.column_dimensions['A'].width = 25
    ws2.column_dimensions['B'].width = 85
    desc_rows = [
        ("字段/项目", "口径与提取标准说明"),
        ("账号与密码", "乐天官方账户账密明文入库，绝不脱敏、绝不生成假密。"),
        ("姓名与假名", "提取自官方購入履歴 Header 及注文詳細お届け先官方原样。"),
        ("其他信息 (G列)", "提取自注文詳細支払い方法或网关，严格标准格式：{卡品牌} **** {后4位}。"),
        ("地址 (J列)", "提取自官方注文詳細お届け先及网关，严格保留 4 行物理原貌（含邮编、县市区门牌、电话）。"),
        ("店铺名称 (K/M/O列)", "乐天官方真实购买店铺名称，执行日期排他规则，杜绝漂移。"),
        ("订单明细 (L/N/P列)", "严格 8 要素 Bana 标准原件订单块，包含注文番号、日時、品名、金额、数量、お届け先、注文者、支払い方法。"),
        ("0 造假铁律", "无订单或未设置门牌/卡号者严格如实留空，绝无任何算法编造或虚拟假数据。")
    ]
    for r_idx, (k, v) in enumerate(desc_rows, 1):
        c1 = ws2.cell(row=r_idx, column=1, value=k)
        c2 = ws2.cell(row=r_idx, column=2, value=v)
        c1.border = thin_border
        c2.border = thin_border
        if r_idx == 1:
            c1.fill = sub_header_fill
            c2.fill = sub_header_fill
            c1.font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
            c2.font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
            c1.alignment = align_center
            c2.alignment = align_center
        else:
            c1.font = Font(name="微软雅黑", size=10, bold=True)
            c2.font = Font(name="微软雅黑", size=10)
            c1.alignment = align_left
            c2.alignment = align_left
        ws2.row_dimensions[r_idx].height = 24

    # --- 编写 Sheet 3: 账号检测与现场快照总览 ---
    s3_headers = ["执行节点", "账号", "登录状态", "官方姓名", "信用卡信息", "收件详细地址", "提取订单数", "全要素验证", "现场快照文件名"]
    for col_idx, h in enumerate(s3_headers, 1):
        cell = ws3.cell(row=1, column=col_idx, value=h)
        cell.font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        cell.fill = sub_header_fill
        cell.alignment = align_center
        cell.border = thin_border
    ws3.row_dimensions[1].height = 28
    
    for r_idx, acc in enumerate(all_accounts, 2):
        ws3.row_dimensions[r_idx].height = 36
        prof = acc.get('profile', {})
        snaps = [Path(s).name for s in acc.get('screenshots', [])]
        snap_str = "\n".join(snaps)
        
        vals = [
            acc.get('node_source', acc.get('worker_tag', 'Worker')),
            acc.get('email', ''),
            acc.get('status', ''),
            prof.get('full_name', ''),
            prof.get('card_info', ''),
            prof.get('address', '').replace('\n', ' | '),
            acc.get('order_count', 0),
            "PASS (100%全要素)" if acc.get('fully_captured') else ("ACTIVE" if acc.get('status') == 'active' else "FAILED"),
            snap_str
        ]
        for col_idx, val in enumerate(vals, 1):
            cell = ws3.cell(row=r_idx, column=col_idx, value=val)
            cell.font = Font(name="微软雅黑", size=10)
            cell.alignment = align_left
            cell.border = thin_border
            
    for c in range(1, 10):
        ws3.column_dimensions[get_column_letter(c)].width = 22
    ws3.column_dimensions['F'].width = 40
    ws3.column_dimensions['I'].width = 35

    # --- 编写 Sheet 4: 订单明细总览 ---
    s4_headers = ["来源节点", "所属账号", "注文番号", "注文日時", "店铺名称", "商品名称", "实付金额(円)", "购买数量"]
    for col_idx, h in enumerate(s4_headers, 1):
        cell = ws4.cell(row=1, column=col_idx, value=h)
        cell.font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border
    ws4.row_dimensions[1].height = 28
    
    s4_row = 2
    for acc in all_accounts:
        node_lbl = acc.get('node_source', acc.get('worker_tag', 'Worker'))
        em = acc.get('email', '')
        for ord_item in acc.get('raw_orders', []):
            ws4.row_dimensions[s4_row].height = 24
            vals = [
                node_lbl,
                em,
                ord_item.get('order_number', ''),
                ord_item.get('order_date', ''),
                ord_item.get('shop_name', ''),
                ord_item.get('item_name', ''),
                ord_item.get('price_yen', ''),
                ord_item.get('quantity', '1')
            ]
            for col_idx, val in enumerate(vals, 1):
                cell = ws4.cell(row=s4_row, column=col_idx, value=val)
                cell.font = Font(name="微软雅黑", size=10)
                cell.alignment = align_left
                cell.border = thin_border
            s4_row += 1
            
    for c in range(1, 9):
        ws4.column_dimensions[get_column_letter(c)].width = 25
    ws4.column_dimensions['F'].width = 45
    
    wb.save(str(out_path))
    total_orders = s4_row - 2
    print(f"🎉 成功生成法定 16 列 Bana 标准真本工作簿: {out_path}")
    print(f"   文件哈希 SHA-256: {sha256_file(out_path)}")
    print(f"   大盘总户数: {len(all_accounts)} | 订单总数: {total_orders}")
    return out_path
